"""Generate synthetic XLSX fixtures with an independent OOXML producer.
Requires openpyxl only for regenerating the checked-in examples.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import json

root = Path(__file__).resolve().parents[1]
out = root / 'examples'
out.mkdir(exist_ok=True)
names = ['调研访谈','原型设计','视觉设计','前端页面','交互联调','数据接入','兼容测试','文档交付','培训支持']
for variant in ['clean','broken','exception','overwritten']:
    wb = Workbook()
    ws = wb.active
    ws.title = '项目报价'
    ws.append(['合成演示数据，不代表真实业务'])
    ws.append(['项目','分类','单位','数量','单价','金额'])
    for r in range(3, 12):
        f = f'=D{r}*E{r}'
        if r == 7:
            if variant == 'broken': f = '=D7*E6'
            if variant == 'exception': f = '=D7*E7*0.9'
            if variant == 'overwritten': f = 4800
        ws.append([names[r-3], '服务', '项', r-1, (r+2)*100, f])
    ws['A13'] = '小计'
    ws['F13'] = '=SUM(F3:F11)'
    ws['A14'] = '含税总计'
    ws['F14'] = '=F13*1.06'
    for cell in ws[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='256045')
    for col in 'ABCDEF': ws.column_dimensions[col].width = 18
    wb.save(out / f'quotation-{variant}.xlsx')

# Values below are derived independently from the fixture input rows, without
# calling the MoonBit parser, translator, or evaluator.
correct = sum((r-1)*(r+2)*100 for r in range(3,12))
broken = correct - 6*900 + 6*800
(out / 'reference-values.json').write_text(json.dumps({
    'provenance':'Synthetic quotation, independently calculated by Python arithmetic',
    'F7': {'before':4800,'after':5400},
    'F13': {'before':broken,'after':correct},
    'F14': {'before':broken*1.06,'after':correct*1.06}
}, ensure_ascii=False, indent=2), encoding='utf8')
print('Generated 4 synthetic XLSX workbooks and independent numeric reference.')
